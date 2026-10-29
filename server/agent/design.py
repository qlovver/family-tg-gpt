"""
design.py — tested helpers for good-looking PowerPoint/Word files.
Copied into every job's work/ folder:  from design import *

Fonts that look right BOTH on iPhone/Mac and in Microsoft Office:
  "Georgia" (elegant serif), "Trebuchet MS" (friendly), "Verdana" (clear),
  "Arial" / "Helvetica Neue", "Times New Roman", "Avenir Next" (iPhone/Mac; Office falls back).
"""

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
import copy

SAFE_FONTS = ["Georgia", "Trebuchet MS", "Verdana", "Arial", "Times New Roman", "Helvetica Neue", "Avenir Next"]

PALETTES = {
    "midnight": ["0F172A", "1E293B", "38BDF8", "F472B6", "F8FAFC"],
    "sunset":   ["2D1E2F", "B33F62", "F9A03F", "F4D35E", "FFF8F1"],
    "forest":   ["17332A", "2F6B4F", "D4A72C", "8FBC94", "F5F8F2"],
    "ocean":    ["0B2545", "13315C", "2EC4B6", "8DA9C4", "F6F9FC"],
    "royal":    ["1C2541", "3A506B", "C9A227", "E0D6C0", "FBF9F4"],
    "candy":    ["3D0C4F", "8E2DE2", "FF6FB5", "FFD166", "FFF7FB"],
    "mono":     ["111111", "333333", "E63946", "BBBBBB", "FFFFFF"],
}


def rgb(hex_):
    hex_ = hex_.lstrip("#")
    return RGBColor(int(hex_[0:2], 16), int(hex_[2:4], 16), int(hex_[4:6], 16))


def new_presentation():
    """16:9 presentation with an empty layout available as prs.slide_layouts[6]."""
    from pptx import Presentation
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_background(slide, color=None, image=None, prs=None):
    """Solid colour background, or a full-bleed picture (e.g. a generated gradient)."""
    if image:
        pic = slide.shapes.add_picture(image, 0, 0, prs.slide_width, prs.slide_height)
        # send to back
        sp_tree = slide.shapes._spTree
        sp_tree.remove(pic._element)
        sp_tree.insert(2, pic._element)
        return pic
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def text(slide, x, y, w, h, content, size=20, color="1F2937", font="Trebuchet MS", bold=False,
         italic=False, align="left", anchor="top", line_spacing=1.1, bullets=False, bullet_color=None, spacing_after=6, accent=None):
    """Text box. content: str (\\n = new paragraph) or list of str / (str, {overrides}).
    **fragment** is bold and, if accent is given, coloured with it."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    items = content.split("\n") if isinstance(content, str) else content
    for i, item in enumerate(items):
        t, ov = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[ov.get("align", align)]
        p.line_spacing = ov.get("line_spacing", line_spacing)
        p.space_after = Pt(ov.get("spacing_after", spacing_after))
        _rich(p, t, size=ov.get("size", size), color=ov.get("color", color), font=ov.get("font", font),
              bold=ov.get("bold", bold), italic=ov.get("italic", italic), accent=ov.get("accent", accent))
        if bullets or ov.get("bullet"):
            _bullet(p, bullet_color or color)
    return box


def _rich(p, t, size, color, font, bold, italic, accent=None):
    """Supports **bold** fragments (rendered bold + accent colour if given)."""
    parts = t.split("**")
    for k, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        f = r.font
        f.size, f.name, f.italic = Pt(size), font, italic
        f.bold = bold or (k % 2 == 1)
        f.color.rgb = rgb(accent if (k % 2 == 1 and accent) else color)


def _bullet(p, color):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.32)))
    pPr.set("indent", str(-Inches(0.28)))
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color.lstrip("#"))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "●")


def shape(slide, kind, x, y, w, h, fill=None, line=None, line_width=1.0, radius=None, alpha=None, shadow=False):
    """kind: 'rect', 'round', 'circle', 'pill', 'triangle', 'arrow', 'chevron', 'hexagon', 'star'."""
    kinds = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "circle": MSO_SHAPE.OVAL,
             "pill": MSO_SHAPE.ROUNDED_RECTANGLE, "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "arrow": MSO_SHAPE.RIGHT_ARROW,
             "chevron": MSO_SHAPE.CHEVRON, "hexagon": MSO_SHAPE.HEXAGON, "star": MSO_SHAPE.STAR_5_POINT}
    s = slide.shapes.add_shape(kinds[kind], Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == "pill":
        s.adjustments[0] = 0.5
    elif kind == "round":
        s.adjustments[0] = radius if radius is not None else 0.12
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            _alpha(s, alpha)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_width)
    else:
        s.line.fill.background()
    if shadow:
        add_shadow(s)
    else:
        _no_shadow(s)
    s.text_frame.text = ""
    return s


def _alpha(s, alpha):
    """alpha 0..1 (1 = opaque)."""
    clr = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    for el in clr.findall(qn("a:alpha")):
        clr.remove(el)
    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 100000)))


def _no_shadow(s):
    spPr = s._element.spPr
    if spPr.find(qn("a:effectLst")) is None:
        etree.SubElement(spPr, qn("a:effectLst"))


def add_shadow(s, blur=0.18, distance=0.06, alpha=0.22, color="000000"):
    """Soft drop shadow (works in PowerPoint, Keynote, LibreOffice)."""
    spPr = s._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(Inches(blur)), dist=str(Inches(distance)),
                          dir="5400000", algn="t", rotWithShape="0")
    c = etree.SubElement(sh, qn("a:srgbClr"), val=color)
    etree.SubElement(c, qn("a:alpha"), val=str(int(alpha * 100000)))
    return s


def gradient(s, colors, angle=90):
    """Linear gradient fill for a shape. colors: list of hex, angle in degrees."""
    s.fill.gradient()
    s.fill.gradient_angle = angle
    stops = s.fill.gradient_stops
    # python-pptx gives 2 stops; set ends, add more via XML if needed
    gs_lst = s.fill._xPr.find(qn("a:gradFill")).find(qn("a:gsLst"))
    for el in list(gs_lst):
        gs_lst.remove(el)
    for i, c in enumerate(colors):
        gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(int(100000 * i / max(1, len(colors) - 1))))
        etree.SubElement(gs, qn("a:srgbClr"), val=c.lstrip("#"))
    return s


def card(slide, x, y, w, h, title, body, accent="2EC4B6", fill="FFFFFF", title_color="1F2937", body_color="4B5563",
         font="Trebuchet MS", number=None, title_size=20, body_size=14, shadow=True):
    """Rounded card with soft shadow, accent bar and optional number badge."""
    shape(slide, "round", x, y, w, h, fill=fill, radius=0.08, shadow=shadow)
    shape(slide, "rect", x + 0.25, y + 0.25, 0.6, 0.07, fill=accent)
    top = y + 0.45
    if number is not None:
        b = shape(slide, "circle", x + 0.25, y + 0.4, 0.6, 0.6, fill=accent)
        b.text_frame.text = str(number)
        p = b.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.runs[0]
        r.font.size, r.font.bold, r.font.name = Pt(18), True, font
        r.font.color.rgb = rgb("FFFFFF")
        top = y + 1.15
    text(slide, x + 0.25, top, w - 0.5, 0.8, title, size=title_size, bold=True, color=title_color, font=font)
    text(slide, x + 0.25, top + 0.75, w - 0.5, h - (top - y) - 0.9, body, size=body_size, color=body_color, font=font, line_spacing=1.2)


def icon_circle(slide, x, y, d, symbol, fill="2EC4B6", color="FFFFFF", font="Segoe UI Symbol"):
    """Circle with a symbol inside (★ ✓ ✦ ➜ ♥ ⚡ ☀ ✎ ⌛ ◆ — or a digit/letter)."""
    c = shape(slide, "circle", x, y, d, d, fill=fill)
    c.text_frame.text = symbol
    p = c.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = p.runs[0]
    r.font.size, r.font.color.rgb = Pt(int(d * 34)), rgb(color)
    return c


TRANSITIONS = {
    "fade": '<p:fade xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
    "push": '<p:push xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" dir="u"/>',
    "wipe": '<p:wipe xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" dir="r"/>',
    "cover": '<p:cover xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" dir="l"/>',
    "split": '<p:split xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" orient="vert" dir="out"/>',
    "zoom": '<p:zoom xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
}


def add_transition(slide, kind="fade", speed="med"):
    """Slide transition (shown in PowerPoint/Keynote when presenting). kind: fade, push, wipe, cover, split, zoom."""
    sld = slide._element
    for el in sld.findall(qn("p:transition")):
        sld.remove(el)
    tr = etree.SubElement(sld, qn("p:transition"), spd=speed)
    tr.append(etree.fromstring(TRANSITIONS.get(kind, TRANSITIONS["fade"])))
    # schema order: cSld, clrMapOvr, transition, timing, ...
    timing = sld.find(qn("p:timing"))
    if timing is not None:
        sld.remove(tr)
        timing.addprevious(tr)
    return slide


def appear_on_click(slide, shapes):
    """Entrance animation: shapes fade in one by one on click (PowerPoint/Keynote)."""
    sld = slide._element
    for el in sld.findall(qn("p:timing")):
        sld.remove(el)
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    pars = []
    n = 3
    for s in shapes:
        sid = s.shape_id
        pars.append(f'''<p:par><p:cTn id="{n}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{n+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{n+2}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{n+3}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{n+4}" dur="500"/><p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>''')
        n += 5
    xml = f'''<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn>
<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
    timing = etree.fromstring(xml)
    ext = sld.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(timing)
    else:
        sld.append(timing)
    return slide


def gradient_image(path, colors, size=(1920, 1080), angle=35, blobs=True, seed=7):
    """Generates a soft gradient background picture (PNG) with optional blurred colour blobs."""
    import numpy as np
    from PIL import Image, ImageDraw, ImageFilter
    import math, random
    w, h = size
    cs = np.array([[int(c.lstrip("#")[i:i + 2], 16) for i in (0, 2, 4)] for c in colors], dtype=float)
    a = math.radians(angle)
    yy, xx = np.mgrid[0:h, 0:w]
    t = xx * math.cos(a) + yy * math.sin(a)
    t = (t - t.min()) / max(1e-9, (t.max() - t.min()))
    t = t * (len(cs) - 1)
    i = np.clip(t.astype(int), 0, max(0, len(cs) - 2))
    f = (t - i)[..., None]
    c1 = cs[i]
    c2 = cs[np.clip(i + 1, 0, len(cs) - 1)]
    arr = (c1 + (c2 - c1) * f).clip(0, 255).astype("uint8")
    img = Image.fromarray(arr, "RGB")
    if blobs:
        random.seed(seed)
        layer = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        d = ImageDraw.Draw(layer)
        for c in list(cs[1:]) + list(cs[:1]):
            r = random.randint(h // 4, h // 2)
            cx, cy = random.randint(0, w), random.randint(0, h)
            d.ellipse((cx - r, cy - r, cx + r, cy + r), fill=tuple(int(v) for v in c) + (90,))
        layer = layer.filter(ImageFilter.GaussianBlur(h // 8))
        img = Image.alpha_composite(img.convert("RGBA"), layer).convert("RGB")
    img.save(path, optimize=True)
    return path


def chart(slide, x, y, w, h, kind, categories, series, colors=None, font="Trebuchet MS", legend=True, labels=False):
    """Native editable chart. kind: 'column', 'bar', 'line', 'pie', 'doughnut', 'area'.
    series: dict name -> list of numbers."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    types = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
             "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA}
    data = CategoryChartData()
    data.categories = categories
    for name, values in series.items():
        data.add_series(name, values)
    gf = slide.shapes.add_chart(types[kind], Inches(x), Inches(y), Inches(w), Inches(h), data)
    ch = gf.chart
    ch.font.name, ch.font.size = font, Pt(12)
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    colors = colors or ["13315C", "2EC4B6", "F9A03F", "B33F62", "8DA9C4", "6B7280"]
    plot = ch.plots[0]
    if kind in ("pie", "doughnut"):
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
    else:
        for i, s in enumerate(plot.series):
            fmt = s.format
            if kind == "line":
                fmt.line.color.rgb = rgb(colors[i % len(colors)])
                fmt.line.width = Pt(3)
            else:
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = rgb(colors[i % len(colors)])
    if labels:
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(11)
    return gf


def check_overflow(prs):
    """Rough check: text boxes whose text likely doesn't fit. Returns a list of warnings."""
    warn = []
    for i, slide in enumerate(prs.slides, 1):
        for s in slide.shapes:
            if not s.has_text_frame or not s.text_frame.text.strip():
                continue
            w_in, h_in = s.width / 914400, s.height / 914400
            lines = 0
            size = 18
            for p in s.text_frame.paragraphs:
                runs = [r for r in p.runs if r.font.size]
                size = runs[0].font.size.pt if runs else size
                chars_per_line = max(1, int(w_in * 72 / (size * 0.55)))
                lines += max(1, -(-len(p.text) // chars_per_line))
            need = lines * size * 1.25 / 72
            if need > h_in * 1.08:
                warn.append(f"слайд {i}: текст «{s.text_frame.text[:30]}…» может не поместиться ({need:.1f} > {h_in:.1f} дюйма)")
    return warn
